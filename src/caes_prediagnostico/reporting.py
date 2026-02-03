from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from caes_prediagnostico.models import ProjectSummary


def build_docx_report(summary: ProjectSummary, output_path: Path) -> None:
    doc = Document()
    title = doc.add_heading(summary.project_name, level=0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    doc.add_paragraph(f"Fecha de generación: {summary.created_at:%Y-%m-%d %H:%M UTC}")

    doc.add_heading("Resumen ejecutivo", level=1)
    doc.add_paragraph(
        "Este informe resume la información recopilada, los cálculos de consumo "
        "y las medidas de ahorro energético con CAEs aplicables."
    )

    doc.add_heading("Totales consolidados", level=1)
    doc.add_paragraph(f"Consumo total: {summary.totals['consumo_kwh']} kWh")
    doc.add_paragraph(f"Coste total: {summary.totals['coste_eur']} €")
    doc.add_paragraph(f"Potencia total: {summary.totals['potencia_kw']} kW")
    doc.add_paragraph(f"Superficie total: {summary.totals['superficie_m2']} m2")

    doc.add_heading("Hallazgos por sección", level=1)
    for section in summary.sections:
        doc.add_heading(section.name, level=2)
        doc.add_paragraph(f"Documentos analizados: {len(section.documents)}")
        if section.extracted_facts:
            for key, values in section.extracted_facts.items():
                doc.add_paragraph(f"{key}: {values}")
        if section.notes:
            doc.add_paragraph("Notas:")
            for note in section.notes:
                doc.add_paragraph(f"- {note}")

    doc.add_heading("Medidas de Ahorro Energético (MAEs)", level=1)
    for medida in summary.measures:
        doc.add_heading(medida["medida"], level=2)
        doc.add_paragraph(f"Ahorro energético: {medida['ahorro_kwh']} kWh")
        doc.add_paragraph(f"Ahorro económico: {medida['ahorro_eur']} €")
        doc.add_paragraph(f"Inversión estimada: {medida['inversion_eur']} €")
        doc.add_paragraph(f"Payback: {medida['payback_anios']} años")

    doc.add_heading("CAEs aplicables", level=1)
    for cae in summary.caes:
        doc.add_paragraph(
            f"Medida: {cae['medida']} | CAE estimado: {cae['cae_estimado_mwh']} MWh"
        )

    doc.add_heading("Presupuesto y ahorros", level=1)
    total_inversion = sum(medida["inversion_eur"] for medida in summary.measures)
    total_ahorro = sum(medida["ahorro_eur"] for medida in summary.measures)
    doc.add_paragraph(f"Inversión total estimada: {total_inversion} €")
    doc.add_paragraph(f"Ahorro anual estimado: {total_ahorro} €")

    doc.add_paragraph(
        "Los cálculos se basan en la información extraída de la documentación "
        "disponible y pueden refinarse con mediciones adicionales."
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def build_pdf_report(summary: ProjectSummary, output_path: Path) -> None:
    styles = getSampleStyleSheet()
    story = []

    story.append(Paragraph(summary.project_name, styles["Title"]))
    story.append(Paragraph(f"Fecha de generación: {summary.created_at:%Y-%m-%d %H:%M UTC}", styles["Normal"]))
    story.append(Spacer(1, 12))

    story.append(Paragraph("Resumen ejecutivo", styles["Heading1"]))
    story.append(
        Paragraph(
            "Este informe resume la información recopilada, los cálculos de consumo "
            "y las medidas de ahorro energético con CAEs aplicables.",
            styles["Normal"],
        )
    )
    story.append(Spacer(1, 12))

    story.append(Paragraph("Totales consolidados", styles["Heading1"]))
    totals_table = Table(
        [
            ["Consumo total (kWh)", summary.totals["consumo_kwh"]],
            ["Coste total (€)", summary.totals["coste_eur"]],
            ["Potencia total (kW)", summary.totals["potencia_kw"]],
            ["Superficie total (m2)", summary.totals["superficie_m2"]],
        ]
    )
    totals_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ]
        )
    )
    story.append(totals_table)
    story.append(Spacer(1, 12))

    story.append(Paragraph("Medidas de Ahorro Energético (MAEs)", styles["Heading1"]))
    for medida in summary.measures:
        story.append(Paragraph(medida["medida"], styles["Heading2"]))
        table = Table(
            [
                ["Ahorro energético (kWh)", medida["ahorro_kwh"]],
                ["Ahorro económico (€)", medida["ahorro_eur"]],
                ["Inversión (€)", medida["inversion_eur"]],
                ["Payback (años)", medida["payback_anios"]],
            ]
        )
        table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey)]))
        story.append(table)
        story.append(Spacer(1, 12))

    story.append(Paragraph("CAEs aplicables", styles["Heading1"]))
    cae_table = Table(
        [["Medida", "CAE estimado (MWh)", "Ahorro ponderado (kWh)"]]
        + [
            [cae["medida"], cae["cae_estimado_mwh"], cae["ahorro_ponderado"]]
            for cae in summary.caes
        ]
    )
    cae_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ]
        )
    )
    story.append(cae_table)
    story.append(Spacer(1, 12))

    total_inversion = sum(medida["inversion_eur"] for medida in summary.measures)
    total_ahorro = sum(medida["ahorro_eur"] for medida in summary.measures)
    story.append(Paragraph("Presupuesto y ahorros", styles["Heading1"]))
    story.append(Paragraph(f"Inversión total estimada: {total_inversion} €", styles["Normal"]))
    story.append(Paragraph(f"Ahorro anual estimado: {total_ahorro} €", styles["Normal"]))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(output_path), pagesize=A4)
    doc.build(story)
