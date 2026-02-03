from __future__ import annotations

import json
import re
from pathlib import Path

import pandas as pd
import pdfplumber
from docx import Document

from caes_prediagnostico.models import SectionDocument, SectionResult

SECTION_NAMES = [
    "facturas",
    "MEE",
    "CAEs",
    "planos",
    "listas_equipos",
    "otros",
]

NUMERIC_PATTERNS = {
    "consumo_kwh": re.compile(r"(?:consumo|energ[ií]a)\s*[:=]\s*([\d.,]+)\s*kwh", re.IGNORECASE),
    "coste_eur": re.compile(r"(?:coste|importe|total)\s*[:=]\s*([\d.,]+)\s*€", re.IGNORECASE),
    "potencia_kw": re.compile(r"(?:potencia|pot\.)\s*[:=]\s*([\d.,]+)\s*kw", re.IGNORECASE),
    "superficie_m2": re.compile(r"(?:superficie|area)\s*[:=]\s*([\d.,]+)\s*m2", re.IGNORECASE),
}


def _load_pdf_text(path: Path) -> str:
    with pdfplumber.open(path) as pdf:
        return "\n".join(page.extract_text() or "" for page in pdf.pages)


def _load_docx_text(path: Path) -> str:
    doc = Document(path)
    return "\n".join(paragraph.text for paragraph in doc.paragraphs)


def _load_excel_text(path: Path) -> str:
    sheets = pd.read_excel(path, sheet_name=None)
    tables = []
    for name, df in sheets.items():
        tables.append(f"[{name}]")
        tables.append(df.to_string(index=False))
    return "\n".join(tables)


def load_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".pdf"}:
        return _load_pdf_text(path)
    if suffix in {".docx"}:
        return _load_docx_text(path)
    if suffix in {".xlsx", ".xls"}:
        return _load_excel_text(path)
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    return ""


def extract_section(section_dir: Path) -> SectionResult:
    result = SectionResult(name=section_dir.name)
    for path in section_dir.rglob("*"):
        if path.is_dir():
            continue
        text = load_text(path)
        if not text:
            result.notes.append(f"Sin texto extraíble: {path.name}")
            continue
        result.documents.append(SectionDocument(path=str(path), text=text))
    _extract_numbers(result)
    return result


def _extract_numbers(result: SectionResult) -> None:
    for doc in result.documents:
        for key, pattern in NUMERIC_PATTERNS.items():
            matches = pattern.findall(doc.text)
            if not matches:
                continue
            numbers = [float(value.replace(".", "").replace(",", ".")) for value in matches]
            result.extracted_facts.setdefault(key, []).extend(numbers)


def consolidate_sections(base_dir: Path) -> list[SectionResult]:
    sections: list[SectionResult] = []
    for name in SECTION_NAMES:
        section_dir = base_dir / name
        if section_dir.exists():
            sections.append(extract_section(section_dir))
    return sections


def export_consolidated_json(sections: list[SectionResult], output_path: Path) -> None:
    payload = [
        {
            "name": section.name,
            "documents": [doc.path for doc in section.documents],
            "extracted_facts": section.extracted_facts,
            "notes": section.notes,
        }
        for section in sections
    ]
    output_path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
